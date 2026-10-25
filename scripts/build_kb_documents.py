"""Regenerates the PDF and DOCX knowledge-base documents from the text below.

Run once:  python scripts/build_kb_documents.py
(The .md and .csv documents are authored directly in data/knowledge_base/.)
All content is fictional and internally consistent with the other KB documents.
"""
from pathlib import Path

from docx import Document as Docx
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors

KB = Path(__file__).resolve().parents[1] / "data" / "knowledge_base"

MANUAL = [
 ("1. Box Contents", [
  "Thank you for choosing the AquaPearl 500 Pro. Before you begin, check that the carton contains everything listed below. If any item is missing or damaged, contact support within 7 days of delivery.",
  "The box contains: the AquaPearl 500 Pro purifier unit with the four filter cartridges and UV lamp pre-installed; one 12-volt power adapter; a wall-mounting bracket with four screws and four wall plugs; a 3-way inlet valve (1/2 inch); 3 metres of 1/4 inch white inlet tubing; 2 metres of red drain tubing; a drain saddle clamp; a chrome faucet with its mounting nut; this user manual; and the Quick Installation Card.",
  "The box does not include a pressure regulator, a booster pump or a pre-filter housing for very muddy water. These are sold separately if your water supply needs them."]),
 ("2. Technical Specifications", [
  "TABLE:Parameter|Value;Purification|5 stages: sediment, carbon block, RO membrane, UV-C, mineral;Pure water output|up to 15 litres per hour;Storage tank|12 litres;Operating inlet pressure|10–60 psi (0.7–4.1 bar);Maximum feed water TDS|2,000 ppm;Feed water temperature|5–38 °C;Power supply|220–240 V AC, 50 Hz via 12 V adapter;Power consumption|36 W while purifying;Dimensions|38 × 26 × 47 cm (W × D × H);Weight|9.2 kg empty",
  "Using the purifier outside these limits may damage it and is not covered by the warranty. If your mains pressure exceeds 60 psi, fit a pressure regulator before the inlet valve."]),
 ("3. Control Panel and Indicator Lights", [
  "The front control panel has a display and three touch buttons: POWER, RESET and WIFI.",
  "The display shows the output TDS in ppm and the remaining life of the filter that is due soonest. A blue LED ring means the UV lamp is working. A steady green LED means the unit is purifying. A blinking red LED means a filter needs attention; the display shows which one. An amber LED indicates a sensor fault.",
  "Press and hold POWER for 3 seconds to switch the unit on or off. Press and hold RESET for 5 seconds to clear a filter alert after replacing the filter. Press and hold WIFI for 3 seconds to enter app pairing mode."]),
 ("4. Operating Instructions", [
  "Open the inlet valve fully and switch on the unit. The pump starts and the tank begins to fill. The first fill takes about 50 minutes.",
  "Before first use you must complete the first-run flushing procedure described in the Installation Guide. Do not drink the water from the first two tanks.",
  "The unit stops automatically when the tank is full and signals with one long beep. In normal use you do not need to switch the unit off. If you will be away for more than 7 days, close the inlet valve and switch the unit off; on return, discard the first full tank."]),
 ("5. Filter Replacement Intervals", [
  "TABLE:Filter|Replace every|SKU;Sediment pre-filter|6 months|AP5-SED;Carbon block|12 months|AP5-CB;RO membrane|24 months|AP5-RO;UV lamp|12 months|AP5-UV;Mineral cartridge|12 months|AP5-MIN",
  "Intervals assume a household of four and feed water below 500 ppm TDS. Harder or muddier water shortens filter life. The app shows the remaining life of each filter in days."]),
 ("6. Filter Maintenance", [
  "To replace the sediment pre-filter: switch off the unit and close the inlet valve. Open the front service door. Rotate the pre-filter cartridge a quarter turn anticlockwise and pull it straight out. Push the new cartridge in until it clicks and rotate it a quarter turn clockwise.",
  "Close the service door, open the inlet valve and switch the unit on. Then press and hold the RESET button for five seconds until the red LED stops blinking. This resets the filter alert and the filter-life counter.",
  "After replacing any filter, flush the system by discarding the first full tank of water, as explained in the Installation Guide. Unusual taste after a filter change is almost always caused by skipping this step."]),
 ("7. Cleaning and Descaling", [
  "Wipe the outside of the unit weekly with a soft damp cloth. Do not use solvents or abrasive cleaners.",
  "Sanitise the storage tank every 3 months: empty the tank, add one sachet of the Pearlz descaling kit (SKU AP5-DESC) dissolved in 1 litre of water, let the unit fill, wait 30 minutes and then discard two full tanks.",
  "Descale the faucet aerator every 6 months by soaking it in white vinegar for 20 minutes. In areas with very hard water, descale every 3 months."]),
 ("8. Safety Warnings", [
  "Do not open the purifier body. There are no user-serviceable parts inside except the filters behind the service door. Opening the body exposes live electrical parts and voids the warranty.",
  "Never look directly at the UV lamp while it is on. Use only the supplied 12 V adapter. Keep the adapter away from water. Do not install the unit outdoors, in direct sunlight, or where it may freeze.",
  "If you notice an electrical smell or water near the power adapter, unplug the unit immediately and contact support."]),
 ("9. Audible Alert Codes", [
  "This table applies to firmware 2.3, which was installed when this manual was printed. Later firmware may change some codes; check the release notes.",
  "TABLE:Sound|Code|Meaning;One long beep|E-01|Tank full (normal);Two short beeps|E-02|Low inlet pressure;Three short beeps|E-03|Sediment pre-filter saturated;Four short beeps|E-04|RO membrane efficiency low;Continuous beeping|E-05|UV lamp fault;Beep every 10 seconds|E-06|Pump overheating;Rapid beeping|E-07|TDS sensor fault;Two long beeps|E-08|Service mode (technicians only)",
  "E-03 is raised when the sediment pre-filter reaches saturation. The red LED blinks together with three short beeps. Replace the pre-filter and hold RESET for five seconds."]),
 ("10. Connecting the Pearlz App", [
  "Install the Pearlz Home app (Android 9 or later, iOS 15 or later). Turn on Bluetooth and location on your phone. Press and hold WIFI on the purifier for 3 seconds until the display shows PAIR, then follow the steps in the app.",
  "The app shows water usage, filter life, TDS and alerts, and lets you order filters and book service visits. The app is optional; the purifier works fully without it."]),
 ("11. Storage and Transport", [
  "If the unit must be moved, close the inlet valve, empty the tank through the faucet, disconnect the tubes and transport it upright. Do not lay the unit on its side, as water remaining in the RO membrane housing can reach the electronics.",
  "For storage longer than one month, remove the RO membrane, seal it in a clean bag with its preservative and store it in a refrigerator."]),
 ("12. Warranty and Service", [
  "The AquaPearl 500 Pro is covered by the Pearlz Warranty Policy. Filter cartridges and the UV lamp are consumables. For the full terms, exclusions and how to claim, see the Warranty Policy and the Service & Maintenance Agreement.",
  "Record your serial number here for reference. The serial number label is inside the front service door."]),
]

INSTALL = [
 ("1. Before You Install", [
  "Check that your site meets these conditions before you start. Inlet water pressure must be between 10 and 60 psi (0.7 to 4.1 bar). Above 60 psi you must fit a pressure regulator; below 10 psi you need a booster pump. Feed water TDS must be below 2,000 ppm.",
  "You need a grounded 220–240 V socket within 1.5 metres of the unit, a cold-water supply line, and a drain pipe for the reject water. Leave at least 10 cm of clearance on each side of the unit and 15 cm above it for filter changes."]),
 ("2. Tools Required", [
  "Drill with 6 mm masonry bit and 12 mm bit (for the faucet hole if your sink has none), screwdriver, adjustable spanner, PTFE thread tape, a tube cutter or sharp blade, a spirit level and a towel. A pencil and measuring tape are useful for marking the bracket position."]),
 ("3. Wall-Mount Installation", [
  "Step 1: Mark the bracket position at least 30 cm above the counter, using the spirit level. Step 2: Drill four 6 mm holes, insert the wall plugs and fix the bracket with the supplied screws. Step 3: Hang the unit on the bracket and check that it sits level.",
  "Step 4: Close the main cold-water supply. Fit the 3-way inlet valve to the cold-water line using PTFE tape. Step 5: Connect the white inlet tube from the valve to the IN port. Step 6: Fit the drain saddle clamp to the drain pipe above the trap and connect the red drain tube. Step 7: Connect the faucet to the OUT port and mount the faucet on the sink."]),
 ("4. Under-Sink Installation", [
  "Place the unit upright on the cabinet floor on a drip tray. Follow steps 4 to 7 of the wall-mount procedure. Make sure no tube is kinked behind the unit and that the drain tube runs downhill without loops.",
  "Under-sink installation is not suitable if the cabinet floor is regularly damp. In that case use the wall-mount method."]),
 ("5. First-Run Flushing Procedure", [
  "Every new purifier and every new filter must be flushed before you drink the water. New carbon and RO cartridges release fine carbon particles and a food-grade preservative. Until these are rinsed out, the water can taste bitter, flat or unusual and may look slightly grey.",
  "To flush: open the inlet valve, switch the unit on, let the tank fill completely (about 50 minutes for a new unit) and then empty it through the faucet. Repeat once, so that you discard the first two full tanks — about 20 minutes of running water after the tank is full each time. After a single filter change, discarding one full tank is usually enough, but discard two if the taste is still unusual.",
  "Skipping this procedure is the most common cause of complaints about taste after installation or after a filter change."]),
 ("6. Professional Installation", [
  "A Pearlz technician can install the unit for PKR 1,500 in Karachi, Lahore, Islamabad and Rawalpindi, and PKR 2,500 in other cities. Book through the Pearlz app or the helpline 0800-PEARLZ. Visits are scheduled within 3 working days in metro cities.",
  "The technician brings a pressure gauge and will fit a regulator or booster pump if needed, which are charged separately."]),
 ("7. Self-Installation and Your Warranty", [
  "You may install the unit yourself without affecting the warranty, provided you follow this guide. The warranty is void if, during self-installation, you drill into or open the purifier body, use fittings or tubes that are not supplied by Pearlz, or connect the unit to inlet pressure above 60 psi without a regulator.",
  "Keep the photos of your installation; support may ask for them if you make a claim."]),
]

REFUND = [
 ("1 Scope", ["This policy explains how returns, replacements and refunds work for products bought from the Pearlz online store. Purchases from retail partners are returned through the partner. Version 2.2, last updated 10 January 2026."]),
 ("2 Return window", ["You may request a return within 14 days of delivery. Requests made after 14 days are handled under the Warranty Policy instead."]),
 ("3 Condition requirements", [
  "3.1 Unopened products in their original sealed packaging receive a full refund of the product price.",
  "3.2 Opened products that have not been installed or used, returned with all original packaging, accessories and manuals, receive a refund minus a 10% restocking fee.",
  "3.3 Products that have been installed or connected to a water supply cannot be returned for a refund. Faults in installed products are handled as warranty repairs or replacements.",
  "3.4 Opened filter cartridges cannot be returned for hygiene reasons."]),
 ("4 Who pays return shipping", [
  "If the product is defective, damaged in transit or not what you ordered, Pearlz arranges and pays for collection. If you are returning a product because you changed your mind, a collection charge of PKR 600 is deducted from the refund."]),
 ("5 Refund processing time", [
  "Refunds are issued after the returned product passes inspection, which takes up to 3 working days after it reaches our warehouse. Refunds then take: 7–10 working days to debit or credit cards; 3–5 working days for bank transfers; 5–7 working days for cash-on-delivery orders, which are refunded by bank transfer; and 2–3 working days to JazzCash or Easypaisa wallets."]),
 ("6 Non-refundable items", [
  "Opened filter cartridges and UV lamps; installation service fees once the technician has visited; descaling kits once opened; and gift cards."]),
 ("7 Return, replacement and warranty repair", [
  "A return gives you your money back within the 14-day window. A replacement is offered when a product is defective on arrival or develops a fault within 7 days of delivery; we exchange it for a new unit of the same model. A warranty repair applies after that, for the warranty period stated in the Warranty Policy. Customers cannot choose a refund instead of a warranty repair after the return window has closed."]),
 ("8 Partial refunds and bundles", [
  "Promotional bundles must be returned complete for a full refund. If only part of a bundle is returned, the refund for the returned items is calculated at their individual list price minus the bundle discount, and cannot exceed the amount paid for the bundle."]),
 ("9 How to request a return", ["Use the Returns page in the app or email returns@pearlzhome.example with your order number and photos of the product and packaging. A return reference is issued within 1 working day."]),
]

PRIVACY = [
 ("1 Who we are", ["Pearlz Home Systems (Pvt.) Ltd. is responsible for personal data collected through our online store, the Pearlz Home app and our support channels. Version 1.4, last updated 1 December 2025."]),
 ("2 Data collected at purchase", ["Name, phone number, email address, delivery address, order history and payment status. We do not store full card numbers; card payments are handled by our payment processor."]),
 ("3 Data collected through the app", [
  "Account details, device model, app version and, if you pair a purifier, water-usage telemetry: litres purified per day, output TDS readings, filter-life counters, alert codes and firmware version. Location permission is used only during Bluetooth pairing and is not stored."]),
 ("4 Why we use it", ["To deliver orders, provide warranty and support, remind you about filter changes, improve product reliability using aggregated telemetry, and meet tax and legal obligations. We do not sell personal data."]),
 ("5 Retention periods", [
  "Order and invoice records are kept for 5 years for tax purposes. App telemetry is kept for 24 months and then deleted or aggregated anonymously. Support chat and email transcripts are kept for 12 months. Marketing preferences are kept until you change them."]),
 ("6 Third parties", [
  "We share only what is necessary with: courier partners (TCS and Leopards Courier) for delivery; our payment processor (PayFast) for card and wallet payments; our cloud hosting provider for storing app data; and authorised service technicians for visits you book."]),
 ("7 Your rights", [
  "You can ask to access, correct, export or delete your personal data. Send requests to privacy@pearlzhome.example from the email address on your account. We respond within 30 days. Deleting your account stops telemetry collection; invoice records required for tax are retained for the legal period."]),
 ("8 Cookies", ["Our website uses essential cookies for the shopping cart and login, and optional analytics cookies that you can refuse in the cookie banner."]),
 ("9 Contact", ["Privacy questions: privacy@pearlzhome.example."]),
]

SERVICE = [
 ("1 Overview", ["This agreement describes Pearlz annual maintenance contracts (AMC) and the procedure for warranty claims. Version 2.0, last updated 15 October 2025."]),
 ("2 Maintenance contract tiers", [
  "2.1 Basic — PKR 4,500 per year: one preventive maintenance visit per year, labour only; parts and filters are billed.",
  "2.2 Plus — PKR 8,900 per year: two visits per year; includes two sediment pre-filters; other parts billed at a 10% discount.",
  "2.3 Premium — PKR 14,500 per year: three visits per year; includes all filters except the RO membrane, and the UV lamp; priority scheduling."]),
 ("3 Response-time SLAs", [
  "For breakdown calls, a technician visit is scheduled within: 5 working days (Basic), 72 hours (Plus), and 24 hours in metro cities or 48 hours elsewhere (Premium). Customers without a contract are scheduled within 7 working days."]),
 ("4 Parts covered and billed", [
  "Parts covered by the product warranty are always free during the warranty period, whatever the contract tier. Outside warranty, the pump, control board and sensors are billed at list price less any tier discount."]),
 ("5 Cancellation", [
  "Contracts cancelled within 30 days of purchase are refunded in full minus the cost of any visit already made. After 30 days, the refund is pro-rata for the unused months minus a 15% administration fee."]),
 ("6 Warranty claim procedure", [
  "6.1 Register the claim through the Pearlz app, the helpline 0800-PEARLZ, or by email to warranty@pearlzhome.example. Provide the invoice, the serial number, a description of the fault and, for leaks or physical faults, photos or a short video.",
  "6.2 You receive a claim reference within 1 working day.",
  "6.3 A technician inspects the unit within 3 working days in metro cities and within 7 working days elsewhere.",
  "6.4 If the fault is covered, the unit is repaired within 10 working days of inspection, or replaced if repair is not possible.",
  "6.5 If the technician finds an excluded cause, such as unauthorised repair or physical damage, the claim is declined in writing and a paid repair quote is offered.",
  "6.6 Claims cover the Pearlz product only. Compensation for damage to other property is not provided under the warranty or this agreement."]),
 ("7 Customer responsibilities", ["Provide safe access to the unit, keep the inlet valve accessible, and inform us of changes to your water supply."]),
]


def build_pdf(path, title, sections, page_breaks=True):
    ss = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=ss["Heading2"], textColor=colors.HexColor("#0E6F6A"))
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=10.5, leading=15)
    story = [Paragraph(title, ss["Title"]), Spacer(1, 6 * mm)]
    for i, (head, paras) in enumerate(sections):
        if page_breaks and i:
            story.append(PageBreak())
        story.append(Paragraph(head, h))
        for p in paras:
            if p.startswith("TABLE:"):
                rows = [r.split("|") for r in p[6:].split(";")]
                t = Table(rows, hAlign="LEFT")
                t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .4, colors.grey),
                                       ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDF4F2")),
                                       ("FONTSIZE", (0, 0), (-1, -1), 9.5)]))
                story += [t, Spacer(1, 4 * mm)]
            else:
                story += [Paragraph(p, body), Spacer(1, 3 * mm)]

    def footer(c, d):
        c.setFont("Helvetica", 8)
        c.drawString(20 * mm, 10 * mm, f"Pearlz Home Systems — {title}")
        c.drawRightString(190 * mm, 10 * mm, f"Page {d.page}")
    SimpleDocTemplate(str(path), pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm).build(
        story, onFirstPage=footer, onLaterPages=footer)


def build_docx(path, title, sections):
    d = Docx()
    d.add_heading(title, 0)
    for head, paras in sections:
        d.add_heading(head, 1)
        for p in paras:
            d.add_paragraph(p)
    d.save(path)


if __name__ == "__main__":
    build_pdf(KB / "product_manual_aquapearl500.pdf", "AquaPearl 500 Pro — User Manual", MANUAL)
    build_pdf(KB / "installation_guide.pdf", "AquaPearl Installation Guide", INSTALL)
    build_docx(KB / "refund_returns_policy.docx", "Refund & Returns Policy", REFUND)
    build_docx(KB / "privacy_policy.docx", "Privacy Policy", PRIVACY)
    build_docx(KB / "service_maintenance_agreement.docx", "Service & Maintenance Agreement", SERVICE)
    print("Built PDF/DOCX knowledge-base documents in", KB)
